"""HR outreach list: the user adds HR contacts (name, email, company), the agent emails each one from a template.

The list lives in its own Excel file (data/hr_contacts.xlsx) so it can be edited in the dashboard or in Excel.
"""
import csv
import io
import logging
import os
import re
import string
import time
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from .activity import friendly_error, record, set_progress, wait_between
from .config import ROOT
from .inbox import check_inbox
from .lock import agent_lock
from .mailer import Mailer
from .sheet import SheetOpen, open_in_spreadsheet as _open_in_spreadsheet

log = logging.getLogger(__name__)

HR_PATH = ROOT / "data" / "hr_contacts.xlsx"
COLUMNS = ["HR Name", "HR Email", "Company", "Role", "Status", "Sent On", "Follow-up Sent On",
           "Replied On", "Message ID", "Notes", "Added On"]
EDITABLE = ["HR Name", "HR Email", "Company", "Role", "Notes"]
# Status: Ready -> Emailed -> Replied | Paused | Failed  (blank = Ready; shown as "Check email" if the address is invalid)
EMAIL_RE = re.compile(r"^[^@\s,;]+@[^@\s,;]+\.[A-Za-z]{2,}$")

DEFAULT_SUBJECT = "Application for {role} role at {company} | {my_name}"
DEFAULT_BODY = """Hi {hr_first_name},

I hope you're doing well. I'm {my_name}, a {role} with {experience} years of experience in {skills}.

I'm exploring new opportunities and would love to be considered for any suitable openings at {company}. I've attached my resume for your reference. My notice period is {notice_period}, and I'd be glad to connect whenever convenient.

Thank you for your time."""

PLACEHOLDERS = {
    "hr_name": "HR's full name",
    "hr_first_name": "HR's first name (or “there”)",
    "company": "Company",
    "role": "Role from the row, else your first target role",
    "my_name": "Your name",
    "experience": "Your years of experience",
    "skills": "Your top skills",
    "notice_period": "Your notice period",
    "phone": "Your mobile",
    "my_email": "Your Gmail",
    "linkedin": "Your LinkedIn URL",
    "naukri": "Your Naukri URL",
}


class TemplateError(ValueError):
    pass


def validate_template(*texts: str) -> None:
    for text in texts:
        try:
            fields = {f for _, f, _, _ in string.Formatter().parse(text) if f is not None}
        except ValueError:
            raise TemplateError("Use { } only around placeholder names, e.g. {company}")
        unknown = fields - set(PLACEHOLDERS)
        if unknown:
            raise TemplateError(f"Unknown placeholder {{{sorted(unknown)[0]}}}. Use one of the listed placeholders.")


def _skills(cfg: dict) -> str:
    from .profile import last_result   # AI-written skill names read better than the search keywords
    result = last_result() or {}
    skills = result.get("key_skills") or cfg["search"]["primary_skills"]
    return ", ".join(skills[:5])


GENERIC_NAMES = {"hr", "h.r", "h.r.", "hr team", "recruiter", "hiring manager", "hiring team", "team", "sir", "madam",
                 "sir/madam", "talent acquisition", "ta", "na", "n/a", "-"}


def values_for(row: dict, cfg: dict) -> dict:
    c = cfg["candidate"]
    name = (row.get("HR Name") or "").strip()
    if name.lower() in GENERIC_NAMES:   # "Hi hr," reads badly; greet them generically instead
        name = ""
    return {
        "hr_name": name or "Hiring Team",
        "hr_first_name": name.split()[0] if name else "there",
        "company": (row.get("Company") or "").strip() or "your company",
        "role": (row.get("Role") or "").strip() or (cfg["search"]["roles"] or ["Software Developer"])[0],
        "my_name": c["name"],
        "experience": c.get("experience_years") or "",
        "skills": _skills(cfg),
        "notice_period": c.get("notice_period") or "flexible",
        "phone": c.get("phone") or "",
        "my_email": cfg["env"]["GMAIL_ADDRESS"],
        "linkedin": c.get("linkedin") or "",
        "naukri": c.get("naukri") or "",
    }


def signature(cfg: dict) -> str:
    c = cfg["candidate"]
    lines = ["Regards,", c["name"]]
    contact = " | ".join(x for x in (c.get("phone"), cfg["env"]["GMAIL_ADDRESS"]) if x)
    if contact:
        lines.append(contact)
    if c.get("linkedin"):
        lines.append(f"LinkedIn: {c['linkedin']}")
    if c.get("naukri"):
        lines.append(f"Naukri: {c['naukri']}")
    return "\n".join(lines)


class _Blank(dict):
    def __missing__(self, key):   # unknown {placeholder}: blank rather than KeyError
        return ""


def _fill(text: str, v: dict) -> str:
    try:
        return text.format_map(_Blank(v))
    except (ValueError, IndexError, AttributeError):   # stray { } or {0}: show the template as typed
        return text


def templates(cfg: dict) -> tuple[str, str]:
    e = cfg["email"]
    return e.get("hr_subject") or DEFAULT_SUBJECT, e.get("hr_body") or DEFAULT_BODY


def render(row: dict, cfg: dict) -> tuple[str, str]:
    v = values_for(row, cfg)
    subject, body = templates(cfg)
    if not (row.get("Company") or "").strip():
        # Without a company, "role at your company" reads oddly in a subject line
        subject = re.sub(r"\s+(?:at|@|with)\s+\{company\}", "", subject)
    return _fill(subject, v).strip(), f"{_fill(body, v).strip()}\n\n{signature(cfg)}"


def follow_up(row: dict, cfg: dict) -> tuple[str, str]:
    v = values_for(row, cfg)
    subject, _ = render(row, cfg)
    sent = str(row.get("Sent On") or "")[:10]
    body = (f"Hi {v['hr_first_name']},\n\nI'm following up on my email" + (f" from {sent}" if sent else "")
            + f" about opportunities at {v['company']}. I'm still very interested and would be glad to connect. "
            "My resume is attached again for reference.\n\nThank you.\n\n" + signature(cfg))
    return f"Re: {subject}", body


HEADER_ALIASES = {
    "HR Name": ("hr name", "name", "hr", "contact name", "recruiter name"),
    "HR Email": ("hr email", "email", "email id", "e-mail", "mail", "recruiter email"),
    "Company": ("company", "company name", "organisation", "organization"),
    "Role": ("role", "position", "designation", "job title", "title"),
}


def open_in_spreadsheet(path: Path | None = None) -> bool:
    return _open_in_spreadsheet(path or HR_PATH)


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d %H:%M")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


class HRSheet:
    """Columns are found by their header text, so rows and columns edited in Excel/LibreOffice keep working."""

    def __init__(self, path: Path | None = None):
        self.path = path = path or HR_PATH
        if path.exists():
            self.wb = load_workbook(path)
            self.ws = self.wb.active
        else:
            self.wb = Workbook()
            self.ws = self.wb.active
            self.ws.title = "HR Contacts"
            self.ws.append(COLUMNS)
            for cell in self.ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="5B4CF0")
            self.ws.freeze_panes = "A2"
            for col, width in zip("ABCDEFGHIJK", [22, 30, 24, 24, 11, 16, 16, 16, 30, 30, 16]):
                self.ws.column_dimensions[col].width = width
        self._map_columns()

    def _map_columns(self) -> None:
        header = [_cell(c.value).lower() for c in self.ws[1]]
        self.col = {}
        for name in COLUMNS:
            names = (name.lower(),) + HEADER_ALIASES.get(name, ())
            idx = next((i for i, h in enumerate(header) if h in names), None)
            if idx is None:   # column missing (older or hand-made sheet): add it at the end
                header.append(name.lower())
                idx = len(header) - 1
                cell = self.ws.cell(row=1, column=idx + 1, value=name)
                cell.font = Font(bold=True)
            self.col[name] = idx + 1

    def rows(self, status: str | None = None):
        for r in range(2, self.ws.max_row + 1):
            row = {name: _cell(self.ws.cell(row=r, column=c).value) for name, c in self.col.items()}
            if not row["HR Email"] and not row["HR Name"]:
                continue
            if not row["Status"]:
                # Rows typed straight into the sheet have no status yet
                row["Status"] = "Ready" if EMAIL_RE.match(row["HR Email"]) else "Check email"
            elif row["Status"] == "Ready" and not EMAIL_RE.match(row["HR Email"]):
                row["Status"] = "Check email"
            if status is None or row["Status"] == status:
                yield r, row

    def emails(self) -> set[str]:
        return {r["HR Email"].lower() for _, r in self.rows()}

    def check(self, row_num: int, expect_email: str | None) -> dict:
        """The row the UI meant to change; fails if rows were added or removed in Excel meanwhile."""
        rows = dict(self.rows())
        row = rows.get(row_num)
        if row is None or (expect_email is not None and row["HR Email"].lower() != expect_email.strip().lower()):
            raise LookupError("The sheet was changed in Excel/LibreOffice. The list has been refreshed; please try again.")
        return row

    def add(self, name: str, email: str, company: str, role: str = "", notes: str = "") -> int:
        email = email.strip()
        if not EMAIL_RE.match(email):
            raise ValueError(f"“{email}” isn't a valid email address")
        if email.lower() in self.emails():
            raise ValueError(f"{email} is already in your HR list")
        r = self.ws.max_row + 1
        values = {"HR Name": name.strip(), "HR Email": email, "Company": company.strip(), "Role": role.strip(),
                  "Status": "Ready", "Notes": notes.strip(), "Added On": datetime.now().strftime("%Y-%m-%d %H:%M")}
        for k, v in values.items():
            self.ws.cell(row=r, column=self.col[k], value=v)
        return r

    def update(self, row_num: int, **fields) -> None:
        for name, value in fields.items():
            self.ws.cell(row=row_num, column=self.col[name], value=value)

    def delete(self, row_num: int) -> None:
        self.ws.delete_rows(row_num)

    def sent_today(self) -> int:
        today = date.today().isoformat()
        return sum(r[c].startswith(today) for _, r in self.rows() for c in ("Sent On", "Follow-up Sent On"))

    def save(self, force: bool = False) -> None:
        """force: write even if the file is open elsewhere (only to record an email that already went out)."""
        if not force and open_in_spreadsheet(self.path):
            raise SheetOpen("hr_contacts.xlsx is open in LibreOffice/Excel. Save and close it there first, "
                            "or make your changes there.")
        tmp = self.path.with_suffix(".tmp.xlsx")
        self.path.parent.mkdir(exist_ok=True)
        self.wb.save(tmp)
        os.replace(tmp, self.path)


def parse_contacts(filename: str, content: bytes) -> list[dict]:
    """Read contacts from an uploaded .xlsx or .csv, or pasted text (one per line: name, email, company, role)."""
    ext = Path(filename).suffix.lower()
    if ext == ".xlsx":
        ws = load_workbook(io.BytesIO(content), read_only=True).active
        table = [[("" if v is None else str(v)) for v in row] for row in ws.iter_rows(values_only=True)]
    else:
        text = content.decode("utf-8-sig", errors="replace")
        table = list(csv.reader(io.StringIO(text), delimiter="\t" if "\t" in text else ","))
    table = [[c.strip() for c in r] for r in table if any(c.strip() for c in r)]
    if not table:
        return []
    header = [h.lower() for h in table[0]]
    find = lambda *names: next((i for i, h in enumerate(header) if any(n in h for n in names)), None)
    idx = {"email": find("email", "mail"), "name": find("name"), "company": find("company", "organisation", "organization"),
           "role": find("role", "position", "title")}
    if idx["email"] is not None:
        body = table[1:]
    else:   # no header row: name, email, company, role, or any order where the email is detectable
        idx, body = {"name": 0, "email": 1, "company": 2, "role": 3}, table
    out = []
    for r in body:
        get = lambda k: r[idx[k]] if idx[k] is not None and idx[k] < len(r) else ""
        email = get("email")
        if not EMAIL_RE.match(email):   # find the email anywhere in the row
            email = next((c for c in r if EMAIL_RE.match(c)), "")
        if email:
            out.append({"name": get("name") if get("name") != email else "", "email": email,
                        "company": get("company"), "role": get("role")})
    return out


def check_replies(cfg: dict) -> int | None:
    """Mark HR rows Replied / Failed (bounced). Gmail is read without holding the sheet lock.
    Returns None if replies couldn't be checked or recorded (follow-ups must wait)."""
    with agent_lock("hr"):
        contacts = {r["HR Email"].lower(): r["Sent On"] for _, r in HRSheet().rows("Emailed") if r["Sent On"]}
    if not contacts:
        return 0
    env = cfg["env"]
    try:
        replies, bounced = check_inbox(env["GMAIL_ADDRESS"], env["GMAIL_APP_PASSWORD"], contacts)
    except Exception as e:
        log.warning("Couldn't check Gmail for HR replies: %s", e)
        return None
    if not (replies or bounced):
        return 0
    if open_in_spreadsheet():
        return None
    with agent_lock("hr"):
        sheet = HRSheet()
        for row_num, r in sheet.rows("Emailed"):
            email = r["HR Email"].lower()
            who = r["HR Name"] or email
            if email in replies:
                sheet.update(row_num, Status="Replied", **{"Replied On": replies[email]})
                record("reply", f"{who} ({r['Company'] or 'HR'}) replied", f"Check your Gmail · {email}")
            elif email in bounced:
                sheet.update(row_num, Status="Failed", Notes=bounced[email])
                record("bounce", f"Email to {who} bounced", bounced[email].removeprefix("Bounced: ").capitalize())
        sheet.save()
    return len(replies)


def _find(sheet: HRSheet, email: str, statuses: tuple[str, ...]) -> tuple[int, dict] | None:
    return next(((n, r) for n, r in sheet.rows() if r["HR Email"].lower() == email and r["Status"] in statuses), None)


def send_hr(cfg: dict, remaining: int, already_emailed: set[str]) -> tuple[int, int]:
    """Email Ready contacts and send due follow-ups, up to `remaining`. Returns (sent, follow-ups).
    The sheet lock is taken per email, so you can keep editing the sheet in the dashboard meanwhile."""
    if open_in_spreadsheet():
        with agent_lock("hr"):
            waiting = any(True for _ in HRSheet().rows("Ready"))
        if waiting:
            record("problem", "Didn't email your HR list", "hr_contacts.xlsx is open in LibreOffice/Excel. Save and close it, then send again.")
        return 0, 0
    try:
        validate_template(*templates(cfg))
    except TemplateError as e:
        record("problem", "Didn't email your HR list", f"Fix your HR email template: {e}")
        return 0, 0
    # Without knowing who replied, a follow-up could go to someone who already answered
    replies_ok = check_replies(cfg) is not None
    ecfg = cfg["email"]
    follow_days = int(ecfg.get("follow_up_days") or 0) if replies_ok else 0
    cutoff = (datetime.now() - timedelta(days=follow_days)).strftime("%Y-%m-%d %H:%M")
    with agent_lock("hr"):
        sheet = HRSheet()
        queue = [("new", r["HR Email"].lower()) for _, r in sheet.rows("Ready")]
        queue += [("follow-up", r["HR Email"].lower()) for _, r in sheet.rows("Emailed")
                  if follow_days and not r["Follow-up Sent On"] and (r["Sent On"] or "9") <= cutoff]
    if remaining <= 0 or not queue:
        return 0, 0
    env = cfg["env"]
    sent = followed = 0
    total = min(len(queue), remaining)
    set_progress(f"Connecting to Gmail to email {total} HR contact{'s' if total != 1 else ''}…")
    with Mailer(env["GMAIL_ADDRESS"], env["GMAIL_APP_PASSWORD"], cfg["candidate"]["resume_path"],
                cfg["candidate"]["name"]) as mailer:
        for i, (kind, email) in enumerate(queue):
            if remaining <= 0:
                break
            emailed_now = False
            with agent_lock("hr"):
                if open_in_spreadsheet():
                    # Without saving "Emailed" we'd email them again next time, so stop instead
                    record("problem", "Paused emailing your HR list", "hr_contacts.xlsx was opened in LibreOffice/Excel. Save and close it, then send again.")
                    break
                sheet = HRSheet()
                found = _find(sheet, email, ("Ready",) if kind == "new" else ("Emailed",))
                if not found or (kind == "follow-up" and found[1]["Follow-up Sent On"]):
                    continue   # paused, deleted or changed since the run started
                row_num, row = found
                to = row["HR Email"]
                who = row["HR Name"] if row["HR Name"] and row["HR Name"].lower() not in GENERIC_NAMES else to
                set_progress(f"HR list: emailing {to} ({sent + followed + 1} of {total})")
                if kind == "new" and email in already_emailed:
                    sheet.update(row_num, Status="Paused", Notes="Already emailed (also in your job list or listed twice)")
                    sheet.save()
                    continue
                try:
                    if kind == "new":
                        subject, body = render(row, cfg)
                        msg_id = mailer.send(to, subject, body)
                        sheet.update(row_num, Status="Emailed", **{"Sent On": datetime.now().strftime("%Y-%m-%d %H:%M"), "Message ID": msg_id})
                        already_emailed.add(email)
                        record("email", f"Emailed {who}" + (f" at {row['Company']}" if row["Company"] else ""), f"To {to} · resume attached")
                        sent += 1
                    else:
                        subject, body = follow_up(row, cfg)
                        mailer.send(to, subject, body, reply_to_id=row["Message ID"] or "")
                        sheet.update(row_num, **{"Follow-up Sent On": datetime.now().strftime("%Y-%m-%d %H:%M")})
                        record("followup", f"Followed up with {who}" + (f" at {row['Company']}" if row["Company"] else ""), f"No reply yet · {to}")
                        followed += 1
                    remaining -= 1
                except Exception as e:
                    if kind == "new":
                        sheet.update(row_num, Status="Failed", Notes=friendly_error(e))
                    record("problem", f"Couldn't email {who}", friendly_error(e))
                    log.warning("HR email to %s failed: %s", to, e)
                try:
                    sheet.save()
                except SheetOpen:
                    # Opened mid-send: the email went out, so record it anyway and stop
                    sheet.save(force=True)
                    record("problem", "Paused emailing your HR list", "hr_contacts.xlsx was opened in LibreOffice/Excel while sending. "
                           "Close it without saving (the agent updated it), then send again.")
                    break
                emailed_now = True
            if emailed_now and remaining > 0 and i < len(queue) - 1:
                wait_between(ecfg["delay_seconds"], f"HR list: sent {sent + followed} of {total}")
    return sent, followed
