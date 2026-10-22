"""Excel tracker: one row per job, deduplicated by job link."""
import os
import re
import subprocess
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

COLUMNS = [
    "Job Link", "Company", "Role", "Recruiter Email", "Mobile", "Location",
    "Profile / JD Summary", "Tech Stack", "Posted Date", "Budget", "Match %",
    "Source", "Status", "Email Subject", "Email Body", "Email Sent On", "Notes",
    "Message ID", "Follow-up Sent On", "Replied On", "Full JD",
]
MAX_CELL = 32000  # Excel's limit per cell is 32,767 characters
# Status values: Drafted -> Approved -> Emailed -> Replied | No Contact | Skipped | Failed
# Statuses meaning we are already in touch with (or about to email) that recruiter
CONTACTING = {"Drafted", "Approved", "Emailed", "Replied", "Applied"}


def same_job_key(job: dict) -> tuple[str, str]:
    norm = lambda s: re.sub(r"[^a-z0-9]", "", str(s or "").lower())
    return norm(job.get("Company")), norm(job.get("Role"))


def split_emails(value) -> list[str]:
    return [e for e in re.split(r"[,;\s]+", str(value or "")) if e]


def first_email(value) -> str:
    return (split_emails(value) or [""])[0].lower()


class SheetOpen(RuntimeError):
    """The sheet is open in LibreOffice/Excel; writing now would lose the user's edits."""


def open_in_spreadsheet(path: Path) -> bool:
    path = Path(path)
    locks = [path.parent / f".~lock.{path.name}#", path.parent / f"~${path.name}"]
    if not any(l.exists() for l in locks):
        return False
    # A lock left behind by a crashed LibreOffice shouldn't block the agent forever
    return subprocess.run(["pgrep", "-f", "soffice"], capture_output=True).returncode == 0


class Tracker:
    def __init__(self, path: str):
        self.path = Path(path)
        if self.path.exists():
            self.wb = load_workbook(self.path)
            self.ws = self.wb.active
            self._add_missing_columns()
        else:
            self.path.parent.mkdir(parents=True, exist_ok=True)
            self.wb = Workbook()
            self.ws = self.wb.active
            self.ws.title = "Jobs"
            self.ws.append(COLUMNS)
            for cell in self.ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="305496")
            self.ws.freeze_panes = "A2"
            for col, width in zip("ABCDEFGHIJKLMNOPQRST", [45, 22, 28, 28, 16, 20, 50, 30, 12, 18, 9, 12, 12, 40, 60, 14, 25, 30, 16, 14]):
                self.ws.column_dimensions[col].width = width
        self.col = {name: i for i, name in enumerate(COLUMNS)}
        self._links = {row[0]: r for r, row in enumerate(self.ws.iter_rows(min_row=2, values_only=True), start=2) if row[0]}

    def _add_missing_columns(self) -> None:
        """Sheets made by older versions lack newer columns; append their headers."""
        header = [c.value for c in self.ws[1]]
        for name in COLUMNS[len(header):]:
            cell = self.ws.cell(row=1, column=len(header) + 1, value=name)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="305496")
            header.append(name)

    def has(self, link: str) -> bool:
        return link in self._links

    def add(self, job: dict) -> None:
        self.ws.append([job.get(c, "") for c in COLUMNS])
        self._links[job["Job Link"]] = self.ws.max_row

    def same_job_keys(self) -> set[tuple[str, str]]:
        """(company, title) of every saved job: the same job is often listed on several sites."""
        return {same_job_key(r) for _, r in self.rows()}

    def row_of(self, link: str) -> int | None:
        return self._links.get(link)

    def rows(self, status: str | None = None):
        """Yield (row_number, dict) for rows, optionally filtered by status."""
        for r, values in enumerate(self.ws.iter_rows(min_row=2, values_only=True), start=2):
            row = dict(zip(COLUMNS, values))
            if status is None or row["Status"] == status:
                yield r, row

    def update(self, row_num: int, **fields) -> None:
        for name, value in fields.items():
            self.ws.cell(row=row_num, column=self.col[name] + 1, value=value)

    def sent_today(self) -> int:
        """Emails sent today, follow-ups included (they count toward the daily limit)."""
        today = date.today().isoformat()
        return sum(str(r[c] or "").startswith(today)
                   for _, r in self.rows() for c in ("Email Sent On", "Follow-up Sent On"))

    def contacting(self) -> dict[str, int]:
        """Recruiter email -> row number, for rows already drafted, approved, emailed or replied."""
        out = {}
        for row_num, r in self.rows():
            email = first_email(r["Recruiter Email"])
            if email and r["Status"] in CONTACTING:
                out.setdefault(email, row_num)
        return out

    def save(self, force: bool = False) -> None:
        """force: write even if the file is open elsewhere (only to record an email that already went out)."""
        if not force and open_in_spreadsheet(self.path):
            raise SheetOpen(f"{self.path.name} is open in LibreOffice/Excel. Save and close it there first.")
        # Write to a temp file and swap it in, so readers never see a half-written sheet
        tmp = self.path.with_suffix(".tmp.xlsx")
        self.wb.save(tmp)
        os.replace(tmp, self.path)
